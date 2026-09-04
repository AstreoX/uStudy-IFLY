"""Extract embedded images from documents (PDF, PPTX, DOCX) for RAG image store."""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass
from typing import Optional

logger = logging.getLogger(__name__)

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None  # type: ignore[assignment]


@dataclass
class ExtractedImage:
    """A single image extracted from a document."""
    image_bytes: bytes
    page_num: Optional[int]
    image_index: int
    ext: str = "png"


class ImageExtractor:
    """Extract images from PDF, PPTX, and DOCX documents."""

    def extract(self, content: bytes, mime_type: str) -> list[ExtractedImage]:
        """Dispatch to the appropriate extractor based on MIME type."""
        if mime_type == "application/pdf":
            return self.extract_from_pdf(content)
        if mime_type in (
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "application/vnd.ms-powerpoint",
        ):
            return self.extract_from_pptx(content)
        if mime_type in (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/msword",
        ):
            return self.extract_from_docx(content)
        return []

    def extract_from_pdf(self, content: bytes) -> list[ExtractedImage]:
        """Extract images from a PDF using PyMuPDF."""
        if fitz is None:
            logger.warning("PyMuPDF (fitz) not installed; skipping PDF image extraction")
            return []

        images: list[ExtractedImage] = []
        try:
            doc = fitz.open(stream=content, filetype="pdf")
            for page_num, page in enumerate(doc, start=1):
                for img_index, img_info in enumerate(page.get_images(full=True)):
                    xref = img_info[0]
                    try:
                        base_image = doc.extract_image(xref)
                        images.append(
                            ExtractedImage(
                                image_bytes=base_image["image"],
                                page_num=page_num,
                                image_index=img_index,
                                ext=base_image.get("ext", "png"),
                            )
                        )
                    except Exception as exc:
                        logger.debug("Failed to extract PDF image xref=%s: %s", xref, exc)
            doc.close()
        except Exception as exc:
            logger.warning("PDF image extraction failed: %s", exc)

        return images

    def extract_from_pptx(self, content: bytes) -> list[ExtractedImage]:
        """Extract images from a PPTX file."""
        try:
            from pptx import Presentation
        except ImportError:
            logger.warning("python-pptx not installed; skipping PPTX image extraction")
            return []

        images: list[ExtractedImage] = []
        try:
            prs = Presentation(io.BytesIO(content))
            for slide_num, slide in enumerate(prs.slides, start=1):
                img_index = 0
                for shape in slide.shapes:
                    if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
                        try:
                            img_bytes = shape.image.blob
                            ext = shape.image.ext or "png"
                            images.append(
                                ExtractedImage(
                                    image_bytes=img_bytes,
                                    page_num=slide_num,
                                    image_index=img_index,
                                    ext=ext,
                                )
                            )
                            img_index += 1
                        except Exception as exc:
                            logger.debug("Failed to extract PPTX image: %s", exc)
        except Exception as exc:
            logger.warning("PPTX image extraction failed: %s", exc)

        return images

    def extract_from_docx(self, content: bytes) -> list[ExtractedImage]:
        """Extract images from a DOCX file (images stored in word/media/)."""
        import zipfile

        images: list[ExtractedImage] = []
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as zf:
                img_index = 0
                for name in zf.namelist():
                    if name.startswith("word/media/"):
                        ext = name.rsplit(".", 1)[-1].lower()
                        if ext in {"png", "jpg", "jpeg", "gif", "bmp", "tiff", "webp"}:
                            images.append(
                                ExtractedImage(
                                    image_bytes=zf.read(name),
                                    page_num=None,
                                    image_index=img_index,
                                    ext=ext,
                                )
                            )
                            img_index += 1
        except Exception as exc:
            logger.warning("DOCX image extraction failed: %s", exc)

        return images
