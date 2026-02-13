"""Word document (DOCX) chunker using python-docx."""

import io
import logging

from docx import Document

from rag.chunking.base import BaseChunker, Chunk

logger = logging.getLogger(__name__)


class DocxChunker(BaseChunker):
    """Word 文档切片器"""

    def chunk(self, content: bytes | str, filename: str | None = None) -> list[Chunk]:
        """
        将 Word 文档切片

        Args:
            content: DOCX 文件内容（字节）
            filename: 文件名（可选）

        Returns:
            切片列表
        """
        if isinstance(content, str):
            raise ValueError("DOCX content must be bytes, not string")

        # 打开 Word 文档
        try:
            doc_stream = io.BytesIO(content)
            doc = Document(doc_stream)
        except Exception as e:
            logger.error("无法打开 Word 文档: %s", e)
            raise ValueError(f"无法解析 Word 文档: {e}")

        # 提取段落和标题
        segments_with_heading: list[tuple[str, str | None]] = []
        current_heading: str | None = None

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            # 检查是否是标题样式
            if para.style and para.style.name and para.style.name.startswith("Heading"):
                current_heading = text
                # 标题本身也作为内容保留
                segments_with_heading.append((text, current_heading))
            else:
                segments_with_heading.append((text, current_heading))

        # 提取表格内容
        for table in doc.tables:
            table_text = self._extract_table_text(table)
            if table_text:
                segments_with_heading.append((table_text, current_heading))

        if not segments_with_heading:
            logger.warning("Word 文档没有可提取的文本内容")
            return []

        # 合并为切片
        chunks = self._merge_with_heading_info(segments_with_heading)

        # 添加元数据
        for chunk in chunks:
            chunk.metadata["source_type"] = "docx"
            if filename:
                chunk.metadata["filename"] = filename

        return chunks

    def _extract_table_text(self, table) -> str:
        """提取表格内容为文本"""
        rows_text = []
        for row in table.rows:
            cells_text = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    cells_text.append(cell_text)
            if cells_text:
                rows_text.append(" | ".join(cells_text))

        if rows_text:
            return "\n".join(rows_text)
        return ""

    def _merge_with_heading_info(
        self, segments_with_heading: list[tuple[str, str | None]]
    ) -> list[Chunk]:
        """
        合并段落为切片，同时记录标题信息

        Args:
            segments_with_heading: (段落文本, 所属标题) 列表

        Returns:
            切片列表，每个切片包含 heading 元数据
        """
        chunks: list[Chunk] = []
        current_content: list[str] = []
        current_tokens = 0
        current_headings: set[str] = set()
        chunk_index = 0

        for segment, heading in segments_with_heading:
            segment_tokens = self.count_tokens(segment)

            # 如果单个段落就超过目标大小，需要进一步分割
            if segment_tokens > self.target_size:
                # 先保存当前累积的内容
                if current_content:
                    chunk_text = "\n".join(current_content)
                    if self.count_tokens(chunk_text) >= self.min_size:
                        chunks.append(
                            Chunk(
                                content=chunk_text,
                                index=chunk_index,
                                token_count=self.count_tokens(chunk_text),
                                metadata={
                                    "headings": list(current_headings - {None}),
                                },
                            )
                        )
                        chunk_index += 1
                    current_content = []
                    current_tokens = 0
                    current_headings = set()

                # 对长段落进行句子级别分割
                sentences = self._split_into_sentences(segment)
                for sentence in sentences:
                    sentence_tokens = self.count_tokens(sentence)

                    if current_tokens + sentence_tokens > self.target_size:
                        if current_content:
                            chunk_text = " ".join(current_content)
                            if self.count_tokens(chunk_text) >= self.min_size:
                                chunks.append(
                                    Chunk(
                                        content=chunk_text,
                                        index=chunk_index,
                                        token_count=self.count_tokens(chunk_text),
                                        metadata={
                                            "headings": list(
                                                current_headings - {None}
                                            ),
                                        },
                                    )
                                )
                                chunk_index += 1

                            overlap_content = self._get_overlap_content(current_content)
                            current_content = overlap_content
                            current_tokens = self.count_tokens(" ".join(current_content))
                            current_headings = {heading} if heading else set()

                    current_content.append(sentence)
                    current_tokens += sentence_tokens
                    if heading:
                        current_headings.add(heading)

            elif current_tokens + segment_tokens > self.target_size:
                # 保存当前切片
                if current_content:
                    chunk_text = "\n".join(current_content)
                    if self.count_tokens(chunk_text) >= self.min_size:
                        chunks.append(
                            Chunk(
                                content=chunk_text,
                                index=chunk_index,
                                token_count=self.count_tokens(chunk_text),
                                metadata={
                                    "headings": list(current_headings - {None}),
                                },
                            )
                        )
                        chunk_index += 1

                    overlap_content = self._get_overlap_content(current_content)
                    current_content = overlap_content + [segment]
                    current_tokens = self.count_tokens("\n".join(current_content))
                    current_headings = {heading} if heading else set()
                else:
                    current_content = [segment]
                    current_tokens = segment_tokens
                    current_headings = {heading} if heading else set()
            else:
                current_content.append(segment)
                current_tokens += segment_tokens
                if heading:
                    current_headings.add(heading)

        # 处理剩余内容
        if current_content:
            chunk_text = "\n".join(current_content)
            if self.count_tokens(chunk_text) >= self.min_size:
                chunks.append(
                    Chunk(
                        content=chunk_text,
                        index=chunk_index,
                        token_count=self.count_tokens(chunk_text),
                        metadata={
                            "headings": list(current_headings - {None}),
                        },
                    )
                )

        return chunks
